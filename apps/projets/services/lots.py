"""Allocation de codes sous verrou du projet et validation Excel sans écritures."""

import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError

from apps.projets.models import Lot
from apps.projets.serializers.lot import LotCreationSerializer

MAX_FICHIER = 5 * 1024 * 1024
MAX_LIGNES = 1000
COLONNES = (
    "nom",
    "mode_execution",
    "type_bordereau",
    "budget_fcfa",
    "date_debut_prevue",
    "date_fin_prevue",
)


def creer_lots(projet, donnees, utilisateur):
    """L'appelant possède transaction.atomic et le verrou select_for_update du projet."""
    existants = list(projet.lots.all())
    codes = {lot.code.upper() for lot in existants}
    ordre = max((lot.ordre for lot in existants), default=0)
    numero = 1
    resultats = []
    for data in donnees:
        while f"L-{numero:02d}" in codes:
            numero += 1
        code = f"L-{numero:02d}"
        ordre += 1
        resultats.append(
            Lot.objects.create(
                projet=projet,
                code=code,
                ordre=ordre,
                cree_par=utilisateur,
                **data,
            )
        )
        codes.add(code)
    return resultats


def normaliser(valeur):
    texte = unicodedata.normalize("NFKD", str(valeur or "").strip().lower())
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return "_".join(texte.replace("(", " ").replace(")", " ").split())


def lire_excel(fichier):
    if not fichier.name.lower().endswith(".xlsx") or not 0 < fichier.size <= MAX_FICHIER:
        raise ValidationError({"fichier": "Fichier .xlsx requis, de 5 Mo maximum."})
    contenu = fichier.read(MAX_FICHIER + 1)
    try:
        with ZipFile(BytesIO(contenu)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024
            ):
                raise ValidationError({"fichier": "Classeur décompressé trop volumineux."})
        classeur = load_workbook(BytesIO(contenu), read_only=True, data_only=False)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError({"fichier": "Classeur Excel invalide."}) from exc
    resultats = []
    erreurs = {}
    try:
        feuille = classeur.active
        feuille.reset_dimensions()
        lignes = feuille.iter_rows(max_row=MAX_LIGNES + 2, max_col=20)
        entetes = [normaliser(cell.value) for cell in next(lignes, [])]
        aliases = {
            "nom_du_lot": "nom",
            "nom_lot": "nom",
            "budget": "budget_fcfa",
            "date_debut": "date_debut_prevue",
            "date_fin": "date_fin_prevue",
        }
        entetes = [aliases.get(nom, nom) for nom in entetes]
        presents = [nom for nom in entetes if nom]
        if "nom" not in presents or len(presents) != len(set(presents)):
            raise ValidationError({"fichier": "Colonne nom requise et colonnes uniques."})
        if set(presents) - set(COLONNES):
            raise ValidationError({"fichier": "Colonnes inconnues ; utilisez le modèle fourni."})
        for numero, ligne in enumerate(lignes, start=2):
            if numero > MAX_LIGNES + 1:
                raise ValidationError({"fichier": "Maximum 1000 lignes de données."})
            if not any(cell.value is not None for cell in ligne):
                continue
            if any(cell.data_type == "f" for cell in ligne):
                erreurs[str(numero)] = {"ligne": "Formules non acceptées."}
                continue
            data = {
                nom: cell.value
                for nom, cell in zip(entetes, ligne, strict=True)
                if nom and cell.value is not None
            }
            data.setdefault("mode_execution", "REGIE")
            data.setdefault("type_bordereau", "FORFAIT")
            for champ in ("mode_execution", "type_bordereau"):
                data[champ] = str(data[champ]).strip().upper()
            for champ in ("date_debut_prevue", "date_fin_prevue"):
                valeur = data.get(champ)
                if isinstance(valeur, datetime):
                    data[champ] = valeur.date().isoformat()
                elif isinstance(valeur, date):
                    data[champ] = valeur.isoformat()
            if "budget_fcfa" in data:
                try:
                    centimes = Decimal(str(data.pop("budget_fcfa"))) * 100
                    if not centimes.is_finite() or centimes != centimes.to_integral_value():
                        raise ValueError
                    data["budget_initial_montant"] = int(centimes)
                except (ValueError, InvalidOperation, OverflowError):
                    erreurs[str(numero)] = {"budget_fcfa": "Montant invalide."}
                    continue
            serializer = LotCreationSerializer(data=data)
            if serializer.is_valid():
                resultats.append(serializer.validated_data)
            else:
                erreurs[str(numero)] = serializer.errors
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError({"fichier": "Contenu du classeur Excel invalide."}) from exc
    finally:
        classeur.close()
    if erreurs:
        raise ValidationError({"lignes": erreurs})
    if not resultats:
        raise ValidationError({"fichier": "Aucun lot à importer."})
    return resultats


def modele_excel():
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Lots"
    feuille.append(COLONNES)
    contenu = BytesIO()
    classeur.save(contenu)
    classeur.close()
    return contenu.getvalue()
