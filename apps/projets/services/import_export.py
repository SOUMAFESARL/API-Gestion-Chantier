"""Le découpage depuis et vers Excel (décision du 30/10/2026).

**Le fichier reprend les colonnes du tableau de démarrage** : une colonne par niveau, dans l'ordre,
puis, facultatives et dans n'importe quel ordre, Unité, Quantité, Montant, Début, Fin. La première
ligne donne le nom des colonnes ; chaque ligne suivante est une tâche, comme dans le tableau.

- **Importer ne crée rien** : le fichier devient un tableau, vérifié comme celui qu'on saisit
  (C-1 à C-7). L'écran le montre ; on l'applique ensuite par la même route que le tableau, à la
  création comme plus tard pour ajouter des branches.
- **Exporter** donne l'arbre au même format, une ligne par tâche : le fichier exporté se
  réimporte tel quel. Les montants n'y figurent que pour qui peut les voir.

Formats : `.xlsx` et `.csv` (séparateur `;` ou `,`, encodage UTF-8 ou Windows-1252).
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from rest_framework.exceptions import ValidationError

from apps.catalogue.models.referentiel import cle_comparaison
from apps.core.enums import UniteMesure
from apps.projets.models import NiveauProjet
from apps.projets.selectors.arbre import charger_arbre
from apps.projets.services import listes_modeles as lm
from apps.projets.services.tableau_demarrage import proposer_tableau, verifier_tableau

__all__ = ["exporter_arbre", "importer_tableau", "lire_fichier"]

TAILLE_MAX = 5 * 1024 * 1024
LIGNES_MAX = 5000

#: Les colonnes de données, reconnues par leur nom (sans accents ni majuscules).
COLONNES_DONNEES = {
    "unite": {"unite", "unites", "u"},
    "quantite": {"quantite", "quantite prevue", "qte", "quantites"},
    "montant": {"montant", "montant (fcfa)", "montant fcfa", "montants", "valeur"},
    "date_debut": {"debut", "date de debut", "date debut", "debut prevu"},
    "date_fin": {"fin", "date de fin", "date fin", "fin prevue"},
}
ENTETES_EXPORT = {
    "unite": "Unité",
    "quantite": "Quantité",
    "montant": "Montant (FCFA)",
    "date_debut": "Début",
    "date_fin": "Fin",
}
UNITES = {
    **{cle_comparaison(v): v for v in UniteMesure.values},
    **{cle_comparaison(label): v for v, label in UniteMesure.choices},
    "m2": "M2",
    "m3": "M3",
    "m": "ML",
    "ml": "ML",
    "kg": "KG",
    "u": "U",
    "unite": "U",
    "ens": "FORFAIT",
    "ff": "FORFAIT",
}


def _cellule(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, date):
        return valeur.isoformat()
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return " ".join(str(valeur).split())


class _PointVirgule(csv.excel):
    delimiter = ";"


def lire_fichier(fichier) -> list[list[str]]:
    """Les lignes du fichier, en texte, lignes vides retirées."""
    nom = (getattr(fichier, "name", "") or "").lower()
    if getattr(fichier, "size", 0) and fichier.size > TAILLE_MAX:
        raise ValidationError({"fichier": "Le fichier dépasse 5 Mo."})
    contenu = fichier.read()
    if nom.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception:
            raise ValidationError({"fichier": "Ce fichier Excel ne s'ouvre pas."}) from None
        feuille = classeur.worksheets[0]
        lignes = [[_cellule(v) for v in ligne] for ligne in feuille.iter_rows(values_only=True)]
    elif nom.endswith(".csv"):
        for encodage in ("utf-8-sig", "cp1252"):
            try:
                texte = contenu.decode(encodage)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t")
        except csv.Error:
            dialecte = _PointVirgule
        lignes = [
            [_cellule(v) for v in ligne] for ligne in csv.reader(io.StringIO(texte), dialecte)
        ]
    else:
        raise ValidationError({"fichier": "Formats acceptés : .xlsx et .csv."})
    lignes = [ligne for ligne in lignes if any(c for c in ligne)]
    if len(lignes) > LIGNES_MAX + 1:
        raise ValidationError({"fichier": f"Au plus {LIGNES_MAX} lignes de tâches."})
    return lignes


def _type_colonne(entete: str) -> str | None:
    cle = cle_comparaison(entete)
    return next((champ for champ, noms in COLONNES_DONNEES.items() if cle in noms), None)


def _nombre(valeur: str):
    if not valeur:
        return None
    propre = valeur.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(propre)
    except InvalidOperation:
        return valeur  # la vérification dira que la valeur est invalide


def _date(valeur: str):
    if not valeur:
        return None
    for format_date in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(valeur[:10], format_date).date().isoformat()
        except ValueError:
            continue
    return valeur


def _niveaux_pour(projet, libelles: list[str]) -> list[dict]:
    """Les niveaux des colonnes : ceux du projet s'il est découpé, sinon ceux du modèle du type
    de projet quand le nom correspond (avec sa liste), sinon un niveau libre."""
    existants = list(NiveauProjet.objects.filter(projet=projet).order_by("ordre"))
    if existants:
        attendus = [cle_comparaison(n.libelle) for n in existants]
        if [cle_comparaison(lib) for lib in libelles] != attendus:
            raise ValidationError(
                {
                    "fichier": "Les colonnes doivent reprendre les niveaux du projet, dans "
                    "l'ordre : " + ", ".join(n.libelle for n in existants) + "."
                }
            )
        return [
            {
                "libelle": n.libelle,
                "liste": lm.reference_liste(n.liste_plateforme or n.liste_entreprise),
                "alerter": n.alerter,
            }
            for n in existants
        ]
    modele = {cle_comparaison(n["libelle"]): n for n in proposer_tableau(projet)["niveaux"]}
    niveaux = []
    for libelle in libelles:
        connu = modele.get(cle_comparaison(libelle))
        niveaux.append(
            {
                "libelle": libelle,
                "liste": connu.get("liste") if connu else None,
                "alerter": bool(connu and connu.get("alerter")),
            }
        )
    return niveaux


def importer_tableau(projet, fichier) -> dict:
    """Le fichier devient un tableau de démarrage, vérifié ; rien n'est écrit."""
    lignes = lire_fichier(fichier)
    if len(lignes) < 2:
        raise ValidationError(
            {"fichier": "Le fichier n'a pas de ligne de tâche sous les en-têtes."}
        )
    entetes = lignes[0]
    colonnes_niveaux, colonnes_donnees = [], {}
    for index, entete in enumerate(entetes):
        champ = _type_colonne(entete) if entete else None
        if champ:
            colonnes_donnees[champ] = index
        elif entete:
            colonnes_niveaux.append(index)
    if not colonnes_niveaux:
        raise ValidationError({"fichier": "Aucune colonne de niveau (Phase, Lot, Activité…)."})
    niveaux = _niveaux_pour(projet, [entetes[i] for i in colonnes_niveaux])
    tableau = []
    for ligne in lignes[1:]:

        def valeur(i, ligne=ligne):
            return ligne[i] if i is not None and i < len(ligne) else ""

        objet = {"cellules": [valeur(i) for i in colonnes_niveaux]}
        unite = valeur(colonnes_donnees.get("unite"))
        if unite:
            objet["unite"] = UNITES.get(cle_comparaison(unite), unite)
        for champ in ("quantite", "montant"):
            nombre = _nombre(valeur(colonnes_donnees.get(champ)))
            if nombre is not None:
                objet[champ] = str(nombre) if isinstance(nombre, Decimal) else nombre
        for champ in ("date_debut", "date_fin"):
            jour = _date(valeur(colonnes_donnees.get(champ)))
            if jour:
                objet[champ] = jour
        tableau.append(objet)
    verdict = verifier_tableau(niveaux, tableau)
    verdict.pop("_lignes", None)
    verdict.pop("lignes_retenues", None)
    return {"niveaux": niveaux, "lignes": tableau, "verification": verdict}


def _lignes_export(projet, montants_visibles: bool) -> tuple[list[str], list[list]]:
    arbre = charger_arbre(projet, actifs_seulement=True)
    niveaux = list(NiveauProjet.objects.filter(projet=projet).order_by("ordre"))
    colonnes = [
        "unite",
        "quantite",
        *(["montant"] if montants_visibles else []),
        "date_debut",
        "date_fin",
    ]
    entetes = [n.libelle for n in niveaux] + [ENTETES_EXPORT[c] for c in colonnes]
    libelles_unites = dict(UniteMesure.choices)
    lignes = []
    numeros = arbre.numerotation()

    def rang(feuille):
        return tuple(int(n) for n in (numeros.get(feuille.pk) or "0").split(".") if n.isdigit())

    for feuille in sorted(arbre.feuilles(), key=rang):
        cellules = arbre.chemin_libelles(feuille)
        cellules += [""] * (len(niveaux) - len(cellules))
        donnees = {
            "unite": libelles_unites.get(feuille.unite, feuille.unite or ""),
            "quantite": float(feuille.quantite_prevue)
            if feuille.quantite_prevue is not None
            else "",
            "montant": feuille.montant if feuille.montant is not None else "",
            "date_debut": feuille.date_debut_prevue or "",
            "date_fin": feuille.date_fin_prevue or "",
        }
        lignes.append(cellules + [donnees[c] for c in colonnes])
    return entetes, lignes


def exporter_arbre(projet, format_fichier: str, *, montants_visibles: bool) -> tuple[bytes, str]:
    """L'arbre au format du tableau : (contenu, type MIME). Le fichier se réimporte tel quel."""
    entetes, lignes = _lignes_export(projet, montants_visibles)
    if format_fichier == "csv":
        tampon = io.StringIO()
        ecrivain = csv.writer(tampon, delimiter=";")
        ecrivain.writerow(entetes)
        for ligne in lignes:
            ecrivain.writerow([c.isoformat() if isinstance(c, date) else c for c in ligne])
        return tampon.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Découpage"
    feuille.append(entetes)
    for ligne in lignes:
        feuille.append(ligne)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return (
        tampon.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
