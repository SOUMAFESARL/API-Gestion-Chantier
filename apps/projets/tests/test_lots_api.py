"""Création multi-lots, accès projet et import Excel transactionnel."""

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Utilisateur
from apps.core.enums import RoleGlobal, RoleProjet
from apps.projets.models import AffectationProjet, Lot, Projet

pytestmark = pytest.mark.django_db


@pytest.fixture
def contexte(schema_demo):
    user = Utilisateur.objects.create_user(
        email="lots.visiteur@demo.ci",
        password="Test12345!",
        nom="Visiteur",
        role_global=RoleGlobal.VISITEUR,
    )
    projet = Projet.objects.create(reference="PRJ-LOTS", nom="Projet", ville="Man")
    AffectationProjet.objects.create(
        projet=projet,
        utilisateur=user,
        role_projet=RoleProjet.VISITEUR,
    )
    client = APIClient(HTTP_HOST="demo.localhost")
    client.force_authenticate(user)
    return client, projet, user


def upload(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    data = BytesIO()
    wb.save(data)
    wb.close()
    return SimpleUploadedFile("lots.xlsx", data.getvalue())


def test_multiple_manual_lots(contexte):
    client, projet, user = contexte
    url = f"/api/v1/projets/{projet.pk}/lots/"
    for numero in (1, 2):
        response = client.post(
            url,
            {
                "nom": "Gros œuvre",
                "mode_execution": "REGIE",
                "type_bordereau": "FORFAIT",
            },
            format="json",
        )
        assert response.status_code == 201, response.data
        assert response.data["code"] == f"L-{numero:02d}"
        assert response.data["projet"] == projet.pk
        assert response.data["budget_initial_montant"] is None
        assert response.data["date_debut_prevue"] is None
    assert len(client.get(url).data) == 2
    assert Lot.objects.filter(projet=projet, cree_par=user).count() == 2


@pytest.mark.parametrize(
    "extra",
    [
        {"budget_initial_montant": -1},
        {"mode_execution": "INCONNU"},
        {"type_bordereau": "INCONNU"},
        {"projet": "INCONNU"},
        {"date_debut_prevue": "2026-10-03", "date_fin_prevue": "2026-10-02"},
    ],
)
def test_invalid_manual_lot(contexte, extra):
    client, projet, _ = contexte
    response = client.post(
        f"/api/v1/projets/{projet.pk}/lots/",
        {
            "nom": "Lot",
            "mode_execution": "REGIE",
            "type_bordereau": "FORFAIT",
            **extra,
        },
        format="json",
    )
    assert response.status_code == 400
    assert not projet.lots.exists()


def test_excel_import_and_atomic_validation(contexte):
    client, projet, _ = contexte
    url = f"/api/v1/projets/{projet.pk}/lots/import/"
    invalid = client.post(
        url,
        {
            "fichier": upload(
                [
                    ["nom", "mode_execution"],
                    ["Lot valide", "REGIE"],
                    ["Lot invalide", "FAUX"],
                ]
            )
        },
        format="multipart",
    )
    assert invalid.status_code == 400, invalid.data
    assert not projet.lots.exists()
    valid = client.post(
        url,
        {
            "fichier": upload(
                [
                    ["nom", "budget_fcfa", "date_debut_prevue"],
                    ["Fondations", 25000000, "2026-10-02"],
                    ["Toiture"],
                ]
            )
        },
        format="multipart",
    )
    assert valid.status_code == 201, valid.data
    assert len(valid.data) == 2
    assert valid.data[0]["budget_initial_montant"] == 2500000000
    assert valid.data[1]["mode_execution"] == "REGIE"
    assert valid.data[1]["type_bordereau"] == "FORFAIT"
    assert valid.data[0]["code"] != valid.data[1]["code"]


def test_excel_invalid_files_and_template(contexte):
    client, projet, _ = contexte
    base = f"/api/v1/projets/{projet.pk}/lots/"
    for fichier in (
        SimpleUploadedFile("fake.xlsx", b"invalid"),
        upload([["nom"], ["=1+1"]]),
        upload([["autre"], ["Lot"]]),
    ):
        assert (
            client.post(base + "import/", {"fichier": fichier}, format="multipart").status_code
            == 400
        )
    response = client.get(base + "modele/")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.active.cell(1, 1).value == "nom"
    wb.close()
    assert not projet.lots.exists()


def test_project_scope_and_inactive_membership(contexte):
    client, projet, user = contexte
    other = Projet.objects.create(reference="PRJ-AUTRE", nom="Autre", ville="Man")
    for suffix in ("", "modele/"):
        assert client.get(f"/api/v1/projets/{other.pk}/lots/{suffix}").status_code == 403
    assert (
        client.post(
            f"/api/v1/projets/{other.pk}/lots/import/",
            {
                "fichier": upload([["nom"], ["Lot"]]),
            },
            format="multipart",
        ).status_code
        == 403
    )
    AffectationProjet.objects.filter(utilisateur=user).update(est_actif=False)
    assert client.get(f"/api/v1/projets/{projet.pk}/lots/").status_code == 403
    client.force_authenticate(None)
    assert client.get(f"/api/v1/projets/{projet.pk}/lots/").status_code in (401, 403)
